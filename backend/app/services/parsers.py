"""Document parsers for PDF, DOCX, TXT, and Markdown formats."""

import logging
import re
from abc import ABC, abstractmethod
from dataclasses import dataclass, field

logger = logging.getLogger(__name__)


@dataclass
class ParsedPage:
    """A single page/section of parsed content."""

    page_number: int
    text: str
    sections: list[str] = field(default_factory=list)


@dataclass
class ParsedDocument:
    """Normalized output from any parser."""

    text: str
    pages: list[ParsedPage] = field(default_factory=list)
    metadata: dict = field(default_factory=dict)
    format: str = "unknown"


class DocumentParser(ABC):
    """Base class for document parsers."""

    @abstractmethod
    def parse(self, file_path: str) -> ParsedDocument:
        """Parse a document file into normalized content."""
        ...

    @staticmethod
    def detect_format(filename: str) -> str:
        """Detect document format from filename."""
        lower = filename.lower()
        if lower.endswith(".pdf"):
            return "pdf"
        if lower.endswith(".docx"):
            return "docx"
        if lower.endswith(".txt"):
            return "txt"
        if lower.endswith(".md") or lower.endswith(".markdown"):
            return "markdown"
        return "unknown"


class PDFParser(DocumentParser):
    """PDF parser using PyMuPDF."""

    def parse(self, file_path: str) -> ParsedDocument:
        import fitz  # PyMuPDF

        logger.info("Parsing PDF: %s", file_path)
        doc = fitz.open(file_path)
        pages: list[ParsedPage] = []
        full_text_parts: list[str] = []

        for i, page in enumerate(doc):
            text = page.get_text()
            page_num = i + 1
            pages.append(ParsedPage(page_number=page_num, text=text.strip()))
            full_text_parts.append(text)

        full_text = "\n\n".join(full_text_parts)
        metadata = {}
        if doc.metadata:
            metadata = {
                "title": doc.metadata.get("title", ""),
                "author": doc.metadata.get("author", ""),
            }

        doc.close()

        return ParsedDocument(
            text=full_text,
            pages=pages,
            metadata=metadata,
            format="pdf",
        )


class DOCXParser(DocumentParser):
    """DOCX parser using python-docx."""

    def parse(self, file_path: str) -> ParsedDocument:
        from docx import Document as DocxDoc

        logger.info("Parsing DOCX: %s", file_path)
        doc = DocxDoc(file_path)
        paragraphs: list[str] = []
        sections: list[str] = []

        for para in doc.paragraphs:
            text = para.text.strip()
            if not text:
                continue
            paragraphs.append(text)
            # Detect headings
            if para.style and para.style.name.startswith("Heading"):
                sections.append(text)

        full_text = "\n\n".join(paragraphs)
        metadata = {}
        if doc.core_properties.title:
            metadata["title"] = doc.core_properties.title
        if doc.core_properties.author:
            metadata["author"] = doc.core_properties.author

        return ParsedDocument(
            text=full_text,
            pages=[ParsedPage(page_number=1, text=full_text, sections=sections)],
            metadata=metadata,
            format="docx",
        )


class TXTParser(DocumentParser):
    """Plain text parser with encoding fallback."""

    def parse(self, file_path: str) -> ParsedDocument:
        logger.info("Parsing TXT: %s", file_path)
        text = self._read_with_encoding_fallback(file_path)
        return ParsedDocument(
            text=text,
            pages=[ParsedPage(page_number=1, text=text)],
            metadata={},
            format="txt",
        )

    @staticmethod
    def _read_with_encoding_fallback(file_path: str) -> str:
        """Read file trying UTF-8 first, then latin-1 fallback."""
        for encoding in ("utf-8", "utf-8-sig", "latin-1"):
            try:
                with open(file_path, encoding=encoding) as f:
                    return f.read()
            except (UnicodeDecodeError, UnicodeError):
                continue
        # Last resort: read as bytes and decode with errors replaced
        with open(file_path, "rb") as f:
            return f.read().decode("utf-8", errors="replace")


class MarkdownParser(DocumentParser):
    """Markdown parser preserving structure."""

    def parse(self, file_path: str) -> ParsedDocument:
        import markdown

        logger.info("Parsing Markdown: %s", file_path)
        raw_text = TXTParser._read_with_encoding_fallback(file_path)

        # Extract headings from raw text
        headings: list[str] = []
        for line in raw_text.split("\n"):
            match = re.match(r"^(#{1,6})\s+(.*)", line)
            if match:
                headings.append(match.group(2).strip())

        # Convert to HTML to verify it parses, but we keep the raw text
        markdown.markdown(raw_text)

        return ParsedDocument(
            text=raw_text,
            pages=[
                ParsedPage(
                    page_number=1,
                    text=raw_text,
                    sections=headings,
                )
            ],
            metadata={"headings": headings},
            format="markdown",
        )


def get_parser(filename: str) -> DocumentParser:
    """Get the appropriate parser for a file format."""
    fmt = DocumentParser.detect_format(filename)
    parsers = {
        "pdf": PDFParser,
        "docx": DOCXParser,
        "txt": TXTParser,
        "markdown": MarkdownParser,
    }
    parser_cls = parsers.get(fmt)
    if parser_cls is None:
        raise ValueError(f"Unsupported document format: {fmt}")
    return parser_cls()
